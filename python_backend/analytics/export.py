import pandas as pd
import pytz
from datetime import timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

MANILA_TZ = pytz.timezone("Asia/Manila")

# --- Threshold constants used in the bottom summary block ---
# Inferred from the values in the source sheet, not from its formulas.
# Confirm against the actual .xls formula bar (Q36/Q38/Q40/U36 etc.)
# before relying on these for Chapter 4.
TOTAL_WAIT_THRESHOLD = timedelta(hours=2.5)      # "Waiting Time ≤/> 2.5 hrs"
EVALUATE_THRESHOLD = timedelta(minutes=30)       # "Evaluate patients ≤/>30 min"
EXAMINE_TREAT_THRESHOLD = timedelta(hours=1)     # "Examine & treat Pts. ≤/>1hr"
CARRYOUT_THRESHOLD = timedelta(minutes=15)       # "Carry out Dr's Orders ≤/>15 min"
DEFAULT_SHIFT_HOURS = 8                          # fallback if a day's span can't be derived

THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=13)


def _to_manila(ts):
    if ts is None or pd.isna(ts):
        return None
    if isinstance(ts, str):
        ts = pd.to_datetime(ts)
    if ts.tzinfo is None:
        ts = ts.tz_localize("UTC")
    return ts.astimezone(MANILA_TZ)


def _fmt_time(ts):
    ts = _to_manila(ts)
    return ts.strftime("%I:%M %p") if ts else ""


def _fmt_hms(td):
    if td is None or pd.isna(td):
        return ""
    total_seconds = int(td.total_seconds())
    if total_seconds < 0:
        return ""
    h, rem = divmod(total_seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}"


def _safe_delta(end, start):
    end, start = _to_manila(end), _to_manila(start)
    if end is None or start is None:
        return None
    return end - start


def _write_merged(ws, cell_range, value, font=None, alignment=CENTER, fill=None):
    ws.merge_cells(cell_range)
    top_left = cell_range.split(":")[0]
    cell = ws[top_left]
    cell.value = value
    cell.alignment = alignment
    if font:
        cell.font = font
    if fill:
        cell.fill = fill


def _sheet_name_for(date_val, clinic_label):
    # Excel sheet names: 31 char max, no : \ / ? * [ ]
    raw = f"{date_val.strftime('%b %d %Y')} {clinic_label}"
    for bad in [":", "\\", "/", "?", "*", "[", "]"]:
        raw = raw.replace(bad, "")
    return raw[:31]


def _write_day_sheet(wb, day_df, date_val, clinic_label):
    ws = wb.create_sheet(title=_sheet_name_for(date_val, clinic_label))

    # --- Title block (10 data columns: A-J) ---
    _write_merged(ws, "A1:J1", "PHILIPPINE HEART CENTER", TITLE_FONT)
    _write_merged(ws, "A2:J2", "OUT-PATIENT DIVISION", BOLD)
    ws.row_dimensions[3].height = 8  # spacer
    _write_merged(ws, "A4:J4", "TIME AND MOTION ANALYSIS", TITLE_FONT)

    ws["A6"] = "Date:"
    ws["A6"].font = BOLD
    ws["B6"] = date_val.strftime("%B %-d, %Y") if hasattr(date_val, "strftime") else str(date_val)
    ws["A7"] = "Clinic:"
    ws["A7"].font = BOLD
    ws["B7"] = clinic_label

    # --- Column headers (row 9-10, two-line like the source) ---
    # NOTE: no Hospital Number column — deliberately excluded (see SECURITY.md
    # open items: anon can SELECT any is_historical=false row, so any
    # hospital-identifier field on live rows would be publicly queryable
    # via the anon key until column-level grants are added).
    header_row1, header_row2 = 9, 10
    headers = [
        ("A", "Queuing", "Time"),
        ("B", "Initial", "Assessment"),
        ("C", "Doctor", "Seen"),
        ("D", "Doctor", "Completed"),
        ("E", "Carry Out", "Completed"),
        ("F", "Queuing Time to", "Initial Assessment"),
        ("G", "Initial Assessment", "to Doctor Seen"),
        ("H", "Doctor Seen", "to Completed"),
        ("I", "Doctor Completed to", "Carry Out Completed"),
        ("J", "Total", "Waiting Time"),
    ]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col, line1, line2 in headers:
        cell = ws[f"{col}{header_row1}"]
        cell.value = f"{line1}\n{line2}"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = CENTER
        cell.border = BORDER
        ws.row_dimensions[header_row1].height = 30

    # --- Data rows ---
    row_idx = header_row2 + 1
    wait_times, evaluate_times, examine_times, carryout_times = [], [], [], []

    for _, r in day_df.sort_values("reg_start").iterrows():
        reg_start, reg_end = r.get("reg_start"), r.get("reg_end")
        consult_start, consult_end = r.get("consult_start"), r.get("consult_end")
        carryout_end = r.get("carryout_end")

        g = _safe_delta(reg_end, reg_start)          # Queuing -> Initial Assessment
        h = _safe_delta(consult_start, reg_end)       # Initial Assessment -> Doctor Seen
        i = _safe_delta(consult_end, consult_start)   # Doctor Seen -> Completed
        j = _safe_delta(carryout_end, consult_end)    # Doctor Completed -> Carry Out
        total = _safe_delta(carryout_end, reg_start)  # Total waiting time (telescopes g+h+i+j)

        # 10 values for 10 headers (A-J) — hospitalNum intentionally omitted
        values = [
            _fmt_time(reg_start),
            _fmt_time(reg_end),
            _fmt_time(consult_start),
            _fmt_time(consult_end),
            _fmt_time(carryout_end),
            _fmt_hms(g), _fmt_hms(h), _fmt_hms(i), _fmt_hms(j), _fmt_hms(total),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")

        if total is not None: wait_times.append(total)
        if g is not None: evaluate_times.append(g)
        if h is not None and i is not None: examine_times.append(h + i)
        if j is not None: carryout_times.append(j)

        row_idx += 1

    for col in "ABCDEFGHIJ":  # 10 letters, matches 10 headers
        ws.column_dimensions[col].width = 13
    ws.freeze_panes = f"A{header_row2 + 1}"

    # --- Summary block ---
    summary_row = row_idx + 2
    n_seen = len(day_df)

    def _bucket(values, threshold):
        le = sum(1 for v in values if v <= threshold)
        gt = sum(1 for v in values if v > threshold)
        return le, gt

    wait_le, wait_gt = _bucket(wait_times, TOTAL_WAIT_THRESHOLD)
    eval_le, eval_gt = _bucket(evaluate_times, EVALUATE_THRESHOLD)
    exam_le, exam_gt = _bucket(examine_times, EXAMINE_TREAT_THRESHOLD)
    carry_le, carry_gt = _bucket(carryout_times, CARRYOUT_THRESHOLD)

    avg_wait = (sum(wait_times, timedelta()) / len(wait_times)) if wait_times else None

    # Doctors on duty isn't derivable from current schema (no per-patient doctor
    # link recorded) — left as a manual override, default 1. See OPEN_ISSUES.md
    # re: "Doctors on Duty" historical recording (Reign to confirm).
    doctors_on_duty = 1

    # FIXED: walrus-on-Series is invalid (ValueError: truth value of a Series
    # is ambiguous). Use .empty checks on the dropna'd Series instead.
    reg_start_notna = day_df["reg_start"].dropna()
    carryout_notna = day_df["carryout_end"].dropna()
    if not reg_start_notna.empty and not carryout_notna.empty:
        span_hours = max(
            (_to_manila(carryout_notna.max()) - _to_manila(reg_start_notna.min())).total_seconds() / 3600,
            1,
        )
    else:
        span_hours = DEFAULT_SHIFT_HOURS

    patient_doctor_ratio = round(n_seen / span_hours) if span_hours else 0

    rows = [
        ("Waiting Time \u2264 2.5 hrs. =", wait_le, "Waiting Time > 2.5 hrs.=", wait_gt,
         "Average Patient's Total Waiting Time =", _fmt_hms(avg_wait)),
        ("Evaluate patients \u226430 min", eval_le, "Evaluate patients >30min", eval_gt,
         "Number of Patient's Seen =", n_seen),
        ("Examine & treat Pts. \u22641hr", exam_le, "Examine & treat Pts. >1hr", exam_gt,
         "Number of Doctors on Duty =", doctors_on_duty),
        ("Carry out Dr's Orders \u226415 min", carry_le, "Carry out Dr's Orders >15 min", carry_gt,
         "Patient to Doctor Ratio Per Hour =", patient_doctor_ratio),
    ]
    for offset, (label1, val1, label2, val2, label3, val3) in enumerate(rows):
        r = summary_row + offset * 2
        _write_merged(ws, f"A{r}:B{r}", label1, BOLD, Alignment(horizontal="left"))
        ws.cell(row=r, column=3, value=val1).alignment = CENTER
        _write_merged(ws, f"D{r}:E{r}", label2, BOLD, Alignment(horizontal="left"))
        ws.cell(row=r, column=6, value=val2).alignment = CENTER
        _write_merged(ws, f"G{r}:I{r}", label3, BOLD, Alignment(horizontal="left"))
        ws.cell(row=r, column=10, value=val3).alignment = CENTER


def build_phc_workbook(df: pd.DataFrame, clinic_label: str = "OPD") -> BytesIO:
    """
    Builds a multi-sheet .xlsx, one sheet per calendar day (Manila time) —
    mirrors PHC's own file structure of one tab per date.

    Expects RAW patients rows (pre-preprocessing.py rename): needs `service`,
    `reg_start`, `reg_end`, `consult_start`, `consult_end`, `carryout_end`.
    Pull straight from Supabase, not through report.py.

    NOTE: rows with reg_start = NULL (any service other than Consultation /
    OPD Screening, per DATABASE_SCHEMA.md) are currently dropped entirely
    during day-grouping below — not yet fixed, flagged separately.
    """
    if df.empty:
        raise ValueError("No rows to export.")

    df = df.copy()
    df["_manila_date"] = df["reg_start"].apply(
        lambda t: _to_manila(t).date() if pd.notna(t) else None
    )
    df = df[df["_manila_date"].notna()]

    if df.empty:
        raise ValueError("No rows with a usable date (reg_start) in this range.")

    wb = Workbook()
    wb.remove(wb.active)  # drop default blank sheet

    for date_val, day_df in sorted(df.groupby("_manila_date")):
        _write_day_sheet(wb, day_df, date_val, clinic_label)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer